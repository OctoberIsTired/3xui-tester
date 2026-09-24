from pathlib import Path

from app.results.writer import ResultStore
from app.security.masking import mask_parameter_values, mask_secrets
from app.parameters.models import ParameterSpec
from app.state.checkpoint import Checkpoint


def test_checkpoint_roundtrip_and_dynamic_result_export(tmp_path: Path) -> None:
    state = Checkpoint("demo", {"xray": "v26.9.9"}, {"abc:1"}, 1)
    path = tmp_path / "state.json"
    state.save(path)
    loaded = Checkpoint.load(path)
    assert loaded and loaded.completed_keys == {"abc:1"} and loaded.failed == 1
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc",
                  "configuration": {"network": "ws", "alpn": ["h2", "http/1.1"]},
                  "result": {"status": "OK", "latency_avg_ms": 12}})
    store.export(["csv", "json", "xlsx"])
    assert "parameter.network" in (tmp_path / "results.csv").read_text(encoding="utf-8")
    assert (tmp_path / "summary.csv").exists() and (tmp_path / "results.json").exists() and (tmp_path / "results.xlsx").exists()
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path / "results.xlsx", read_only=False)["results"]
    rows = sheet.iter_rows()
    headers = [cell.value for cell in next(rows)]
    row = [cell.value for cell in next(rows)]
    assert row[headers.index("parameter.alpn")] == '["h2","http/1.1"]'
    assert sheet.freeze_panes == "G2"


def test_secret_masking() -> None:
    assert mask_secrets({"api_token": "very-secret", "nested": {"password": "x"}}) == {"api_token": "***", "nested": {"password": "***"}}
    spec = ParameterSpec.from_dict("custom", {"type": "string", "values": ["sensitive"],
                                              "path": ["streamSettings", "realitySettings", "privateKey"]})
    assert mask_parameter_values({"custom": "sensitive"}, [spec]) == {"custom": "***"}


def test_skipped_candidate_is_not_counted_as_failed_measurement(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "skip", "configuration": {},
                  "result": {"status": "SKIPPED", "reason_code": "reality_target_or_sni_missing"}})
    store.append({"test_id": 2, "run": 1, "configuration_hash": "ok", "configuration": {},
                  "result": {"status": "OK", "score": 1.0}})
    summaries = {item["configuration_hash"]: item for item in store._summaries(store.records())}
    assert summaries["skip"]["skipped_tests"] == 1
    assert summaries["skip"]["failed_tests"] == 0
    assert summaries["ok"]["success_rate"] == 1.0


def test_xlsx_only_export_does_not_create_csv_and_accepts_structured_values(tmp_path: Path) -> None:
    directory = tmp_path / "xlsx-only"
    store = ResultStore(directory)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc",
                  "configuration": {"dest_override": ["http", "tls", "quic", "fakedns"]},
                  "result": {"status": "OK", "latency_avg_ms": 10.25}})
    store.export(["xlsx"])
    assert (directory / "results.xlsx").exists()
    assert not (directory / "results.csv").exists()
    assert not (directory / "summary.csv").exists()
    assert not (directory / "results.json").exists()


def test_xlsx_flattens_nested_metric_groups_into_columns(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc", "configuration": {},
                  "result": {"status": "FAILED", "stage": "screening",
                             "screening": {"success_rate": 0.0, "tcp_connect_ok": True},
                             "gate_failures": ["success_rate<1"]}})
    store.export(["xlsx"])
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path / "results.xlsx")["results"]
    headers = [cell.value for cell in sheet[1]]
    assert "screening.success_rate" in headers
    assert "screening.tcp_connect_ok" in headers
    assert sheet.cell(2, headers.index("gate_failures") + 1).value == '["success_rate<1"]'


def test_summary_and_xlsx_dashboard_report_repeat_stability(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    for run, p95, speed in ((1, 100.0, 120.0), (2, 120.0, 140.0)):
        store.append({"test_id": 1, "run": run, "configuration_hash": "abc",
                      "configuration": {"tls_alpn": ["h2"]},
                      "result": {"status": "OK", "score": 0.9, "latency_p95_ms": p95,
                                 "download_mbps": speed}})
    summary = store._summaries(store.records())
    assert summary[0]["latency_p95_stddev"] > 0
    assert summary[0]["download_stddev"] > 0
    store.export(["xlsx"])
    from openpyxl import load_workbook
    workbook = load_workbook(tmp_path / "results.xlsx")
    assert workbook.sheetnames[:2] == ["Dashboard", "Candidates"]
    assert len(workbook["Dashboard"]._charts) == 4
    candidates = workbook["Candidates"]
    assert candidates["A2"].value == "#1"
    assert candidates["B2"].value == '{"tls_alpn": ["h2"]}'
    assert candidates["C2"].value == 1
    assert candidates["E2"].value == "abc"
    assert candidates["M2"].value > 0
    assert workbook["Dashboard"]["B43"].value == "#1"
    assert workbook["Dashboard"]["C43"].value == candidates["B2"].value
    score, latency, scatter, stability = workbook["Dashboard"]._charts
    assert score.series[0].val.numRef.f == "'Candidates'!$J$2"
    assert latency.series[0].val.numRef.f == "'Candidates'!$L$2"
    assert stability.series[0].val.numRef.f == "'Candidates'!$M$2"
    assert scatter.series[0].xVal.numRef.f == "'Candidates'!$L$2"
    assert scatter.series[0].yVal.numRef.f == "'Candidates'!$N$2"


def test_xlsx_chart_numbers_map_to_complete_configurations(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    for test_id, network, latency, speed in ((7, "ws", 95, 120), (8, "tcp", 110, 90)):
        store.append({"test_id": test_id, "run": 1, "configuration_hash": network,
                      "configuration": {"network": network, "tls_alpn": ["h2"]},
                      "result": {"status": "OK", "score": 0.8,
                                 "latency_p95_ms": latency, "download_mbps": speed}})
    store.export(["xlsx"])
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "results.xlsx")
    rows = list(workbook["Candidates"].values)
    assert [row[0] for row in rows[1:]] == ["#1", "#2"]
    assert [row[2] for row in rows[1:]] == [7, 8]
    assert '"network": "ws"' in rows[1][1]
    assert '"network": "tcp"' in rows[2][1]
    assert [workbook["Dashboard"].cell(row, 2).value for row in (43, 44)] == ["#1", "#2"]
    scatter = workbook["Dashboard"]._charts[2]
    assert len(scatter.series) == 2
    assert [series.title.strRef.strCache.pt[0].v if series.title.strRef else series.title.v
            for series in scatter.series] == ["#1", "#2"]
