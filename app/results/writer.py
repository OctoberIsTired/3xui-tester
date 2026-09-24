from __future__ import annotations

import csv
import json
import statistics
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from app.security.masking import mask_secrets


class ResultStore:
    def __init__(self, directory: Path):
        self.directory = directory
        self.directory.mkdir(parents=True, exist_ok=True)
        self.results_path = directory / "results.jsonl"
        self.errors_path = directory / "errors.jsonl"

    def append(self, record: dict[str, Any]) -> None:
        with self.results_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(mask_secrets(record), ensure_ascii=False, sort_keys=True) + "\n")
            handle.flush()

    def error(self, record: dict[str, Any]) -> None:
        with self.errors_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(mask_secrets(record), ensure_ascii=False, sort_keys=True) + "\n")
            handle.flush()

    def records(self) -> list[dict[str, Any]]:
        if not self.results_path.exists():
            return []
        return [json.loads(line) for line in self.results_path.read_text(encoding="utf-8").splitlines() if line.strip()]

    def export(self, formats: Iterable[str]) -> None:
        records = self.records()
        requested = {item.lower() for item in formats}
        if "csv" in requested:
            self._write_csv(records, self.directory / "results.csv", flatten=True)
            self._write_csv(self._summaries(records), self.directory / "summary.csv", flatten=False)
        if "json" in requested:
            # JSONL remains the durable journal; this is the optional convenient array.
            (self.directory / "results.json").write_text(
                json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        if "xlsx" in requested:
            self._write_xlsx(records)

    def _flatten(self, record: dict[str, Any]) -> dict[str, Any]:
        result = {"phase": record.get("phase", "search"), "test_id": record.get("test_id"),
                  "candidate_rank": record.get("candidate_rank"),
                  "configuration_hash": record.get("configuration_hash"), "run": record.get("run"),
                  "status": record.get("result", {}).get("status")}
        result.update({f"parameter.{key}": value for key, value in record.get("configuration", {}).items()})
        for key, value in record.get("result", {}).items():
            if key == "status":
                continue
            if isinstance(value, dict):
                result.update(_flatten_mapping(key, value))
            else:
                result[key] = value
        return result

    def _write_csv(self, records: list[dict[str, Any]], path: Path, *, flatten: bool) -> None:
        rows = [self._flatten(record) for record in records] if flatten else records
        fields = list(dict.fromkeys(field for row in rows for field in row))
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    def _summaries(self, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
        groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        for record in records:
            groups[(record.get("phase", "search"), record["configuration_hash"])].append(record)
        summaries = []
        for (phase, digest), group in groups.items():
            metrics = [item.get("result", {}) for item in group]
            success = [item for item in metrics if item.get("status") == "OK"]
            skipped = [item for item in metrics if item.get("status") == "SKIPPED"]
            attempted = len(group) - len(skipped)
            summary = {"phase": phase, "test_id": group[0].get("test_id"),
                       "candidate_rank": group[0].get("candidate_rank"),
                       "configuration_hash": digest, "configuration": json.dumps(group[0]["configuration"], ensure_ascii=False),
                       "tests": len(group), "successful_tests": len(success),
                       "failed_tests": attempted - len(success), "skipped_tests": len(skipped),
                       "success_rate": len(success) / attempted if attempted else 0.0}
            for output, source in (("score", "score"), ("connect_time", "connect_time_ms"),
                                   ("latency", "latency_avg_ms"), ("latency_p95", "latency_p95_ms"),
                                   ("download", "download_mbps")):
                values = [float(item[source]) for item in metrics if item.get(source) is not None]
                if values:
                    summary.update({f"{output}_min": min(values), f"{output}_avg": sum(values) / len(values), f"{output}_max": max(values)})
                    if len(values) > 1:
                        summary[f"{output}_stddev"] = statistics.stdev(values)
            summaries.append(summary)
        summaries.sort(key=lambda item: (
            0 if item["phase"] == "validation" else 1,
            item.get("candidate_rank") if item.get("candidate_rank") is not None else 10**9,
            -float(item.get("score_avg", 0)),
        ))
        return summaries

    def _write_xlsx(self, records: list[dict[str, Any]]) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table, TableStyleInfo
        except ImportError as error:
            raise RuntimeError("XLSX export requires openpyxl; install project dependencies") from error
        workbook = Workbook()
        summaries = self._summaries(records)
        validation_records = [item for item in records if item.get("phase") == "validation"]
        sources: list[tuple[str, list[dict[str, Any]], bool]] = [("results", records, True)]
        if validation_records:
            sources.append(("validation", validation_records, True))
        sources.append(("summary", summaries, False))
        if validation_records:
            sources.append(("validation_summary", [item for item in summaries if item["phase"] == "validation"], False))
        for index, (title, source, flatten) in enumerate(sources):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = title
            rows = [self._flatten(item) for item in source] if flatten else source
            fields = list(dict.fromkeys(field for row in rows for field in row))
            sheet.append(fields)
            for row in rows:
                sheet.append([_xlsx_value(row.get(field)) for field in fields])
            # The detailed sheet is very wide; keep identifiers visible while scrolling.
            sheet.freeze_panes = "G2" if flatten else "A2"
            sheet.sheet_view.showGridLines = False
            sheet.auto_filter.ref = sheet.dimensions
            header_fill = PatternFill("solid", fgColor="1F4E78")
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.row_dimensions[1].height = 24
            for column_index, field in enumerate(fields, start=1):
                letter = get_column_letter(column_index)
                values = [field, *(_xlsx_value(row.get(field)) for row in rows[:200])]
                width = min(60, max(10, max(len(str(value)) if value is not None else 0 for value in values) + 2))
                sheet.column_dimensions[letter].width = width
                if field == "success_rate":
                    for cell in sheet[letter][1:]:
                        cell.number_format = "0.0%"
                elif field.endswith("_ms") or field.endswith("_mbps"):
                    for cell in sheet[letter][1:]:
                        cell.number_format = "0.00"
            if rows and fields:
                table = Table(displayName=f"{title.title()}Table", ref=sheet.dimensions)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2", showFirstColumn=False,
                    showLastColumn=False, showRowStripes=True, showColumnStripes=False,
                )
                sheet.add_table(table)
        self._add_dashboard(workbook, summaries)
        workbook.save(self.directory / "results.xlsx")

    @staticmethod
    def _add_dashboard(workbook: Any, summaries: list[dict[str, Any]]) -> None:
        """Add a compact decision view without replacing detailed audit sheets."""
        from openpyxl.chart import BarChart, Reference, ScatterChart, Series
        from openpyxl.chart.label import DataLabelList
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        dashboard = workbook.create_sheet("Dashboard", 0)
        candidates = workbook.create_sheet("Candidates", 1)
        dashboard.sheet_view.showGridLines = candidates.sheet_view.showGridLines = False
        navy, blue, cyan, green, amber, white, muted = (
            "11263A", "1B3E5B", "37C6F4", "35D09A", "F7C948", "F8FAFC", "B6C2D0"
        )
        thin = Side(style="thin", color="2D5876")
        candidates.append([
            "№ на графике", "Параметры конфигурации", "ID теста", "Фаза", "Хеш",
            "Запуски", "Успешные", "Неудачные", "Доля успеха", "Средний балл", "σ балла",
            "Средний p95 (мс)", "σ p95 (мс)", "Средняя скорость (Мбит/с)",
            "σ скорости (Мбит/с)", "Статус",
        ])
        for number, summary in enumerate(summaries, start=1):
            candidates.append([
                f"#{number}", summary.get("configuration"), summary.get("test_id"),
                summary.get("phase"), summary.get("configuration_hash"),
                summary.get("tests", 0), summary.get("successful_tests", 0), summary.get("failed_tests", 0),
                summary.get("success_rate", 0), summary.get("score_avg"), summary.get("score_stddev"),
                summary.get("latency_p95_avg"), summary.get("latency_p95_stddev"),
                summary.get("download_avg"), summary.get("download_stddev"),
                ("SKIPPED" if summary.get("skipped_tests", 0) == summary.get("tests", 0) else
                 "OK" if summary.get("successful_tests", 0) else "FAILED"),
            ])
        for cell in candidates[1]:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="Arial", size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        candidates.freeze_panes = "C2"
        candidates.auto_filter.ref = candidates.dimensions
        candidates.row_dimensions[1].height = 24
        for index, width in enumerate((16, 60, 12, 14, 24, 10, 12, 12, 14, 14, 12,
                                       18, 16, 26, 24, 12), start=1):
            candidates.column_dimensions[get_column_letter(index)].width = width
        for row in candidates.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")
            row[1].alignment = Alignment(vertical="center", wrap_text=True)
        for column in (9,):
            for cell in candidates.iter_cols(min_col=column, max_col=column, min_row=2):
                for value in cell:
                    value.number_format = "0.0%"
        for column in (10, 11):
            for cell in candidates.iter_cols(min_col=column, max_col=column, min_row=2):
                for value in cell:
                    value.number_format = "0.000"
        for column in (12, 13, 14, 15):
            for cell in candidates.iter_cols(min_col=column, max_col=column, min_row=2):
                for value in cell:
                    value.number_format = "0.0"
        if summaries:
            table = Table(displayName="CandidatesTable", ref=f"A1:P{len(summaries) + 1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                                  showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            candidates.add_table(table)

        dashboard.merge_cells("B2:Q2")
        dashboard["B2"] = "3xui-tester — результаты эксперимента"
        dashboard.merge_cells("B3:Q3")
        dashboard["B3"] = "Номер #N на графиках и в рейтинге соответствует строке #N на листе Candidates; там указаны все параметры."
        dashboard.merge_cells("B4:Q4")
        dashboard["B4"] = "Кандидаты с одним прогоном не получают оценку стабильности: запустите не менее 5 повторов для уверенного сравнения."
        for row in dashboard["B2:Q4"]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=navy)
        dashboard["B2"].font = Font(name="Arial", size=16, bold=True, color=white)
        for cell in (dashboard["B3"], dashboard["B4"]):
            cell.font = Font(name="Arial", size=10, italic=True, color=muted)
        for row in range(2, 5):
            dashboard.row_dimensions[row].height = 22

        record_count = sum(int(item.get("tests", 0)) - int(item.get("skipped_tests", 0)) for item in summaries)
        success_count = sum(int(item.get("successful_tests", 0)) for item in summaries)
        best_score = max(summaries, key=lambda item: float(item.get("score_avg", 0)), default=None)
        successful = [item for item in summaries if item.get("latency_p95_avg") is not None]
        best_latency = min(successful, key=lambda item: float(item["latency_p95_avg"]), default=None)
        fastest = max((item for item in summaries if item.get("download_avg") is not None),
                      key=lambda item: float(item["download_avg"]), default=None)
        def numbered_label(summary: dict[str, Any] | None) -> str:
            if summary is None:
                return "нет данных"
            number = next(index for index, item in enumerate(summaries, start=1) if item is summary)
            return f"#{number} · {_configuration_label(summary.get('configuration'), summary.get('configuration_hash'))}"
        cards = [
            ("B5:E8", "УСПЕШНЫЕ ЗАПУСКИ", success_count / record_count if record_count else 0, "0%", cyan,
             f"{success_count} из {record_count} запусков"),
            ("F5:I8", "ЛУЧШИЙ SCORE", best_score.get("score_avg") if best_score else None, "0.000", green,
             numbered_label(best_score)),
            ("J5:M8", "МИНИМАЛЬНЫЙ p95", best_latency.get("latency_p95_avg") if best_latency else None, "0.0\" ms\"", cyan,
             numbered_label(best_latency)),
            ("N5:Q8", "МАКС. СКОРОСТЬ", fastest.get("download_avg") if fastest else None, "0.0\" Mbps\"", amber,
             numbered_label(fastest)),
        ]
        for range_ref, label, value, number_format, color, subtitle in cards:
            first, last = range_ref.split(":")
            start_column = dashboard[first].column_letter
            end_column = dashboard[last].column_letter
            for row in dashboard[range_ref]:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="152E44")
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            dashboard.merge_cells(f"{start_column}5:{end_column}5")
            dashboard.merge_cells(f"{start_column}6:{end_column}7")
            dashboard.merge_cells(f"{start_column}8:{end_column}8")
            dashboard[f"{start_column}5"] = label
            dashboard[f"{start_column}5"].font = Font(name="Arial", size=9, bold=True, color=muted)
            dashboard[f"{start_column}6"] = value
            dashboard[f"{start_column}6"].number_format = number_format
            dashboard[f"{start_column}6"].font = Font(name="Arial", size=20, bold=True, color=color)
            dashboard[f"{start_column}6"].alignment = Alignment(vertical="center")
            dashboard[f"{start_column}8"] = subtitle
            dashboard[f"{start_column}8"].font = Font(name="Arial", size=9, italic=True, color=white)

        for title, target in (("SCORE ПО КОНФИГУРАЦИЯМ", "B10:I10"), ("p95 LATENCY (ms)", "J10:Q10"),
                              ("СКОРОСТЬ И p95", "B25:I25"), ("СТАБИЛЬНОСТЬ p95", "J25:Q25")):
            dashboard.merge_cells(target)
            cell = dashboard[target.split(":")[0]]
            cell.value = title
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="Arial", size=10, bold=True, color=white)
        if summaries:
            max_row = len(summaries) + 1
            categories = Reference(candidates, min_col=1, min_row=2, max_row=max_row)
            score_chart = BarChart()
            score_chart.title = "Средний балл по конфигурациям"
            score_chart.x_axis.title = "Номер конфигурации (#)"
            score_chart.y_axis.title = "Средний балл (0–1)"
            score_chart.y_axis.numFmt = "0.000"
            score_chart.height, score_chart.width = 8.0, 14.5
            score_chart.add_data(Reference(candidates, min_col=10, min_row=1, max_row=max_row), titles_from_data=True)
            score_chart.set_categories(categories)
            score_chart.dLbls = DataLabelList(showVal=True)
            score_chart.legend = None
            dashboard.add_chart(score_chart, "B11")
            latency_chart = BarChart()
            latency_chart.title = "Средняя задержка p95"
            latency_chart.x_axis.title = "Номер конфигурации (#)"
            latency_chart.y_axis.title = "Задержка p95 (мс)"
            latency_chart.y_axis.numFmt = "0.0"
            latency_chart.height, latency_chart.width = 8.0, 14.5
            latency_chart.add_data(Reference(candidates, min_col=12, min_row=1, max_row=max_row), titles_from_data=True)
            latency_chart.set_categories(categories)
            latency_chart.dLbls = DataLabelList(showVal=True)
            latency_chart.legend = None
            dashboard.add_chart(latency_chart, "J11")
            scatter = ScatterChart()
            scatter.title = "Скорость и задержка p95"
            scatter.x_axis.title = "Задержка p95 (мс)"
            scatter.y_axis.title = "Скорость загрузки (Мбит/с)"
            scatter.x_axis.numFmt = "0.0"
            scatter.y_axis.numFmt = "0.0"
            scatter.height, scatter.width = 8.0, 14.5
            for number, summary in enumerate(summaries, start=1):
                if summary.get("latency_p95_avg") is None or summary.get("download_avg") is None:
                    continue
                row = number + 1
                x_values = Reference(candidates, min_col=12, min_row=row, max_row=row)
                y_values = Reference(candidates, min_col=14, min_row=row, max_row=row)
                scatter.series.append(Series(y_values, x_values, title=f"#{number}"))
            scatter.dLbls = DataLabelList(showSerName=True)
            scatter.legend = None
            dashboard.add_chart(scatter, "B26")
            stability_chart = BarChart()
            stability_chart.title = "Разброс задержки p95"
            stability_chart.x_axis.title = "Номер конфигурации (#)"
            stability_chart.y_axis.title = "Стандартное отклонение p95 (мс)"
            stability_chart.y_axis.numFmt = "0.0"
            stability_chart.height, stability_chart.width = 8.0, 14.5
            stability_chart.add_data(Reference(candidates, min_col=13, min_row=1, max_row=max_row), titles_from_data=True)
            stability_chart.set_categories(categories)
            stability_chart.dLbls = DataLabelList(showVal=True)
            stability_chart.legend = None
            dashboard.add_chart(stability_chart, "J26")

        dashboard.merge_cells("B41:Q41")
        dashboard["B41"] = "РЕЙТИНГ И СТАБИЛЬНОСТЬ КАНДИДАТОВ"
        dashboard["B41"].fill = PatternFill("solid", fgColor=blue)
        dashboard["B41"].font = Font(name="Arial", size=10, bold=True, color=white)
        headers = ["№", "Параметры конфигурации", "Запуски", "Успех", "Балл", "p95, мс",
                   "σ p95, мс", "Скорость, Мбит/с", "σ скорости"]
        for column, header in enumerate(headers, start=2):
            cell = dashboard.cell(42, column, header)
            cell.fill = PatternFill("solid", fgColor="244965")
            cell.font = Font(name="Arial", size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal="center")
        for number, (row_number, summary) in enumerate(enumerate(summaries, start=43), start=1):
            dashboard.cell(row_number, 2, f"#{number}")
            dashboard.cell(row_number, 3, summary.get("configuration"))
            dashboard.cell(row_number, 4, summary.get("tests", 0))
            dashboard.cell(row_number, 5, summary.get("success_rate", 0))
            dashboard.cell(row_number, 6, summary.get("score_avg"))
            dashboard.cell(row_number, 7, summary.get("latency_p95_avg"))
            dashboard.cell(row_number, 8, summary.get("latency_p95_stddev"))
            dashboard.cell(row_number, 9, summary.get("download_avg"))
            dashboard.cell(row_number, 10, summary.get("download_stddev"))
        for row in dashboard.iter_rows(min_row=42, max_row=max(43, 42 + len(summaries)), min_col=2, max_col=10):
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if cell.row > 42:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
                    cell.font = Font(name="Arial", size=10, color="172033")
                    cell.alignment = Alignment(vertical="center")
        for row in range(43, 43 + len(summaries)):
            dashboard.cell(row, 5).number_format = "0.0%"
            dashboard.cell(row, 6).number_format = "0.000"
            for column in (7, 8, 9, 10):
                dashboard.cell(row, column).number_format = "0.0"
        for column, width in {"B": 10, "C": 40, "D": 11, "E": 11, "F": 12,
                              "G": 12, "H": 14, "I": 20, "J": 16}.items():
            dashboard.column_dimensions[column].width = width
        for column in range(11, 18):
            dashboard.column_dimensions[get_column_letter(column)].width = 12
        for row in range(5, 9):
            dashboard.row_dimensions[row].height = 22


def _configuration_label(serialized: Any, digest: Any) -> str:
    try:
        config = json.loads(serialized) if isinstance(serialized, str) else dict(serialized or {})
    except (TypeError, ValueError):
        config = {}
    if config:
        return "; ".join(f"{key}={_xlsx_value(value)}" for key, value in sorted(config.items()))
    return f"базовая конфигурация ({str(digest or 'unknown')[:12]})"


def _xlsx_value(value: Any) -> Any:
    """Convert structured Xray parameter values into portable spreadsheet cells."""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return value


def _flatten_mapping(prefix: str, value: dict[str, Any]) -> dict[str, Any]:
    """Give nested metric groups real spreadsheet columns instead of JSON blobs."""
    flattened: dict[str, Any] = {}
    for key, item in value.items():
        name = f"{prefix}.{key}"
        if isinstance(item, dict):
            flattened.update(_flatten_mapping(name, item))
        else:
            flattened[name] = item
    return flattened
